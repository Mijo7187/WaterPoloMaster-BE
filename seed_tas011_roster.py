"""
One-off roster import — PK Taš011 (club id 5) from the season 25/26 xlsx.

Reads `excel/PK-Tas011-Korisnici-Grupe (2).xlsx` (sheet "Korisnici", 181 players)
and seeds two axes:

  * ACADEMY-level, shared  (company_id = club.academy_id)
      - 4 selections   Obuka / Mlađi takmičari / Takmičari / Rekreacija
      - 4 memberships  SWIMMING + MONTHLY, one per selection (7000/7500/8500/6500)
      - 2 seasons      2025/2026 (past) and 2026/2027 (current)
      - 8 groups       (season x selection)

  * CLUB-level, billing    (company_id = 5)
      - 181 users + their USER wallets
      - 330 group_user rows (everyone in 25/26, the September signers in 26/27)
      - 149 MEMBERSHIP contracts   (only where the Sep 26 cell is not blank)
      - 149 September-2026 installments — and nothing later; October onward
        belongs to app.scheduler.billing_jobs._monthly_recurring_billing_job
      - 1282 payments              (145 September + 1137 historical)

The xlsx encodes each month cell by BOTH value and fill colour, and the two are
read and cross-checked against each other:

    no fill   + empty      -> not enrolled that month, nothing written
    C6E0B4    + positive   -> paid that amount        -> COMPLETED payment
    F8CBAD    + 0          -> enrolled, did not pay   -> DEBT payment

A cell whose colour and value disagree is a hard error naming the cell — the
encoding is the whole contract with this file and must not be guessed at.

Everything is get-or-create by natural key and the whole run is ONE transaction,
so `--dry-run` executes every insert and every assertion and then rolls back.
Note that a dry run still burns primary-key sequence values; `nextval` is not
transactional. That is cosmetic — ids simply start higher on the real run.

Known limitations, all deliberate:
  * Users are matched on (normalised name tokens, date_of_birth). 16 rows carry
    the placeholder birthdate 1900-01-01, so correcting one of those in the UI
    and re-running this script would create a duplicate rather than match.
  * Rows where the roster put the surname in the Ime column are stored as
    written. Matching is order-insensitive; the stored names are not corrected.
  * Scholarship contracts are ACTIVE at amount 0. `_validate_contract_update`
    rejects amount <= 0, so those 4 cannot be edited through the API.
  * DEBT payments are terminal in the API today: `_ALLOWED_TRANSITIONS` in
    payment_service.py has no entry for DEBT, so the 308 debt rows this writes
    cannot be moved to COMPLETED until that mapping gains `DEBT: {COMPLETED}`.

Run with:  PYTHONIOENCODING=utf-8 python seed_tas011_roster.py --dry-run
           PYTHONIOENCODING=utf-8 python seed_tas011_roster.py
"""

import argparse
import os
import sys
import unicodedata
from collections import OrderedDict
from datetime import date, datetime
from decimal import Decimal
from zoneinfo import ZoneInfo

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

# Import every model so the SQLAlchemy mapper can resolve all relationships.
import app.features.users.users_models  # noqa: F401,E402
import app.features.company.company_model  # noqa: F401,E402
import app.features.sifarnici.country.country_model  # noqa: F401,E402
import app.features.sifarnici.city.city_model  # noqa: F401,E402
import app.features.sifarnici.exercise_option.exercise_option_model  # noqa: F401,E402
import app.features.sifarnici.expense_category.expense_category_model  # noqa: F401,E402
import app.features.sifarnici.income_category.income_category_model  # noqa: F401,E402
import app.features.sifarnici.selection.selection_model  # noqa: F401,E402
import app.features.group.group_model  # noqa: F401,E402
import app.features.group_user.group_user_model  # noqa: F401,E402
import app.features.training.training_model  # noqa: F401,E402
import app.features.training_users.training_users_model  # noqa: F401,E402
import app.features.training_segments.training_segments_model  # noqa: F401,E402
import app.features.tournament.tournament_model  # noqa: F401,E402
import app.features.tournament_users.tournament_users_model  # noqa: F401,E402
import app.features.season.season_model  # noqa: F401,E402
import app.features.membership.membership_model  # noqa: F401,E402
import app.features.contract.contract_model  # noqa: F401,E402
import app.features.contract_installment.contract_installment_model  # noqa: F401,E402
import app.features.wallet.wallet_model  # noqa: F401,E402
import app.features.payment.payment_model  # noqa: F401,E402

from app.core.config import settings  # noqa: E402
from app.core.db.database import SessionLocal  # noqa: E402
from app.core.security import hash_password  # noqa: E402
from app.features.company.company_model import Company  # noqa: E402
from app.features.contract.contract_model import (  # noqa: E402
    Contract,
    ContractStatus,
    ContractType,
)
from app.features.contract.contract_service import (  # noqa: E402
    compute_contract_status,
    month_bounds,
)
from app.features.contract_installment.contract_installment_model import (  # noqa: E402
    ContractInstallment,
)
from app.features.group.group_model import Group  # noqa: E402
from app.features.group_user.group_user_model import GroupUser  # noqa: E402
from app.features.membership.membership_model import (  # noqa: E402
    BillingType,
    Membership,
    Program,
)
from app.features.payment.payment_model import (  # noqa: E402
    PayableType,
    Payment,
    PaymentStatus,
    PaymentTypeCode,
)
from app.features.season.season_model import Season  # noqa: E402
from app.features.sifarnici.selection.selection_model import Selection  # noqa: E402
from app.features.users.users_models import User, UserRole  # noqa: E402
from app.features.wallet.wallet_model import Wallet, WalletOwnerType  # noqa: E402
from app.utils.dateUtils import club_today  # noqa: E402


# ============================================
# CONSTANTS
# ============================================

CLUB_ID = 5
XLSX_PATH = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "excel",
    "PK-Tas011-Korisnici-Grupe (2).xlsx",
)
SHEET_NAME = "Korisnici"
FIRST_DATA_ROW = 2

CLUB_TZ = ZoneInfo(settings.CLUB_TIMEZONE)

# Canonical selection names. Every write goes through these constants —
# UNIQUE(company_id, name) is byte-exact, so "Mlađi Takmičari" and
# "Mlađi takmičari" would be two different rows.
SEL_OBUKA = "Obuka"
SEL_MLADJI = "Mlađi takmičari"
SEL_TAKMICARI = "Takmičari"
SEL_REKREACIJA = "Rekreacija"

SELECTION_ORDER = [SEL_OBUKA, SEL_MLADJI, SEL_TAKMICARI, SEL_REKREACIJA]
SELECTION_PRICES = {
    SEL_OBUKA: 7000,
    SEL_MLADJI: 7500,
    SEL_TAKMICARI: 8500,
    SEL_REKREACIJA: 6500,
}

# Roster "Grupa" labels, keyed by normalise_label() output.
GRUPA_TO_SELECTION = {
    "obuka": SEL_OBUKA,
    "mladi takmicari": SEL_MLADJI,
    "takmicari": SEL_TAKMICARI,
    "rekreacija": SEL_REKREACIJA,
}

SEASON_PREV = ("2025/2026", date(2025, 9, 1), date(2026, 8, 31), False)
SEASON_CURR = ("2026/2027", date(2026, 9, 1), date(2027, 8, 31), True)

# Columns I..T, in sheet order. The last entry is September 2026 — the only
# month that gets a contract and an installment.
MONTH_COLUMNS = [
    (9, (2025, 10)),
    (10, (2025, 11)),
    (11, (2025, 12)),
    (12, (2026, 1)),
    (13, (2026, 2)),
    (14, (2026, 3)),
    (15, (2026, 4)),
    (16, (2026, 5)),
    (17, (2026, 6)),
    (18, (2026, 7)),
    (19, (2026, 8)),
    (20, (2026, 9)),
]
SEPTEMBER_INDEX = len(MONTH_COLUMNS) - 1
CONTRACT_START = date(2026, 9, 1)

COL_IME = 2
COL_PREZIME = 3
COL_DOB = 4
COL_GRUPA = 5
COL_CENA = 7

# Fill colours, compared on the last 6 hex digits (the file writes a "00"
# alpha prefix rather than "FF").
FILL_PAID = "C6E0B4"     # green
FILL_DEBT = "F8CBAD"     # salmon, the legend's "crveno"
FILL_ZERO_PRICE = "FCE4D6"  # orange, on the Mesečna cena column

PROGRAM = Program.SWIMMING
BILLING = BillingType.MONTHLY

EMAIL_DOMAIN = "tas011.local"
DEFAULT_PHONE = "000000000"
DEFAULT_PASSWORD = "Tas011!Import"
PLACEHOLDER_DOB = date(1900, 1, 1)
# Not [PLAYER]: UserCreate requires position + default_team alongside the
# PLAYER role and the roster has neither, so PLAYER would seed rows the API
# itself could never have created. Squad membership is group_user's job.
ROSTER_ROLES = [UserRole.USER.value]

DESCRIPTION_FMT = "Članarina {:02d}/{:d}"

STATE_BLANK = "BLANK"
STATE_PAID = "PAID"
STATE_DEBT = "DEBT"


# ============================================
# NORMALISATION
# ============================================

# đ/Đ have NO canonical decomposition, so NFD leaves them alone and they would
# survive the combining-mark filter. Translate them first.
_STROKE = str.maketrans({"đ": "d", "Đ": "D", "ð": "d", "Ð": "D"})


def strip_diacritics(value: str) -> str:
    """č ć š ž -> c c s z, and đ -> d. Returns ASCII for Latin-Serbian input."""
    value = value.translate(_STROKE)
    value = unicodedata.normalize("NFD", value)
    return "".join(ch for ch in value if not unicodedata.combining(ch))


def normalise_label(value: str) -> str:
    """Fold a free-text label (a Grupa cell) down to a lookup key."""
    return " ".join(strip_diacritics(value).casefold().split())


def name_key(first: str, last: str, dob: date):
    """Match key for a person: sorted name tokens + birthdate.

    Sorting the tokens makes the key insensitive to which column the surname
    landed in — the roster has rows like Ime="Prljević", Prezime="Luka".
    """
    tokens = strip_diacritics(f"{first} {last}").casefold()
    tokens = "".join(ch if ch.isalnum() or ch.isspace() else " " for ch in tokens)
    return (tuple(sorted(tokens.split())), dob)


def email_slug(value: str) -> str:
    value = strip_diacritics(value).casefold()
    return "".join(ch for ch in value if ch.isalnum())


def email_for(first: str, last: str, row_no: int, taken: set) -> str:
    """A deterministic placeholder address on a non-routable domain.

    Collisions take the roster's own row number rather than a running counter,
    so the same row always produces the same address across runs.
    """
    base = f"{email_slug(first)}.{email_slug(last)}" or f"igrac{row_no}"
    candidate = f"{base}@{EMAIL_DOMAIN}"
    if candidate in taken:
        candidate = f"{base}.{row_no}@{EMAIL_DOMAIN}"
    if candidate in taken:
        raise RuntimeError(f"Could not build a unique email for row {row_no}.")
    taken.add(candidate)
    return candidate


# ============================================
# COUNTERS
# ============================================


class Counters:
    """created / skipped per entity, printed as a table at the end."""

    def __init__(self, verbose=False):
        self.rows = OrderedDict()
        self.verbose = verbose
        self.warnings = []

    def note(self, tag, created, message=None, loud=False):
        row = self.rows.setdefault(tag, {"created": 0, "skipped": 0})
        row["created" if created else "skipped"] += 1
        if message and (loud or self.verbose):
            print(f"  [{tag:<11}] {'created' if created else 'exists '} {message}")

    def warn(self, message):
        self.warnings.append(message)
        print(f"  [WARN       ] {message}")

    def print_summary(self):
        print("\n  entity                  created   existing")
        print("  " + "-" * 42)
        for tag, row in self.rows.items():
            print(f"  {tag:<22} {row['created']:>7}   {row['skipped']:>8}")
        if self.warnings:
            print(f"\n  {len(self.warnings)} warning(s) above.")


# ============================================
# XLSX PARSING
# ============================================


def _rgb6(cell):
    """The cell's fill colour as 6 hex digits, or None when it has no fill."""
    fill = cell.fill
    if fill is None or fill.patternType != "solid":
        return None
    colour = fill.start_color
    if colour is None or colour.type != "rgb" or not colour.rgb:
        return None
    return str(colour.rgb)[-6:].upper()


def _month_state(cell, sheet_ref):
    """(state, amount) for one month cell, with value and colour cross-checked."""
    fill = _rgb6(cell)
    value = cell.value

    if fill is None and value in (None, ""):
        return STATE_BLANK, None
    if fill == FILL_PAID and isinstance(value, (int, float)) and value > 0:
        return STATE_PAID, int(value)
    if fill == FILL_DEBT and isinstance(value, (int, float)) and value == 0:
        return STATE_DEBT, 0

    raise ValueError(
        f"{sheet_ref}: cannot read the cell — fill={fill!r} value={value!r}. "
        "Expected no-fill+empty (not enrolled), "
        f"{FILL_PAID}+positive (paid) or {FILL_DEBT}+0 (unpaid)."
    )


def parse_roster(path):
    """Read the sheet into plain dicts before anything touches the database."""
    from openpyxl import load_workbook

    if not os.path.exists(path):
        raise SystemExit(f"Roster file not found: {path}")

    workbook = load_workbook(path, data_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        raise SystemExit(
            f"Sheet {SHEET_NAME!r} not in {path} (found: {workbook.sheetnames})."
        )
    sheet = workbook[SHEET_NAME]

    rows = []
    for excel_row in range(FIRST_DATA_ROW, sheet.max_row + 1):
        first = sheet.cell(excel_row, COL_IME).value
        last = sheet.cell(excel_row, COL_PREZIME).value
        if first in (None, "") and last in (None, ""):
            continue

        dob_cell = sheet.cell(excel_row, COL_DOB).value
        if isinstance(dob_cell, datetime):
            dob = dob_cell.date()
        elif isinstance(dob_cell, date):
            dob = dob_cell
        else:
            dob = PLACEHOLDER_DOB

        grupa_raw = str(sheet.cell(excel_row, COL_GRUPA).value or "").strip()
        selection_name = GRUPA_TO_SELECTION.get(normalise_label(grupa_raw))
        if selection_name is None:
            raise ValueError(
                f"Row {excel_row}: unknown Grupa {grupa_raw!r}. "
                f"Known: {sorted(GRUPA_TO_SELECTION)}"
            )

        price_cell = sheet.cell(excel_row, COL_CENA)
        price = price_cell.value
        if not isinstance(price, (int, float)):
            raise ValueError(f"Row {excel_row}: Mesečna cena is {price!r}, not a number.")
        price = int(price)
        # The legend says orange marks a zero price; hold the file to it.
        orange = _rgb6(price_cell) == FILL_ZERO_PRICE
        if (price == 0) != orange:
            raise ValueError(
                f"Row {excel_row}: Mesečna cena {price} disagrees with its fill "
                f"({_rgb6(price_cell)!r}); orange must mean 0 and only 0."
            )

        months = []
        for column, (year, month) in MONTH_COLUMNS:
            cell = sheet.cell(excel_row, column)
            months.append(_month_state(cell, f"{cell.coordinate} (row {excel_row})"))

        rows.append({
            "row_no": excel_row,
            "first_name": str(first or "").strip(),
            "last_name": str(last or "").strip(),
            "dob": dob,
            "selection_name": selection_name,
            "price": price,
            "months": months,
        })

    return rows


def validate_roster(roster):
    """Two guards that must hold before a single row is written."""
    seen = {}
    for row in roster:
        key = name_key(row["first_name"], row["last_name"], row["dob"])
        if key in seen:
            raise ValueError(
                f"Rows {seen[key]} and {row['row_no']} share the match key {key} — "
                "they would collapse into one user. Resolve in the sheet first."
            )
        seen[key] = row["row_no"]

        joined = strip_diacritics(f"{row['first_name']} {row['last_name']}")
        if not joined.isascii():
            raise ValueError(
                f"Row {row['row_no']}: {joined!r} is not ASCII after folding — "
                "a non-Latin character would make the match key unstable."
            )


def expected_counts(roster):
    """Everything the run should produce, derived from the roster itself."""
    signers = [r for r in roster if r["months"][SEPTEMBER_INDEX][0] != STATE_BLANK]
    sep_completed = sum(
        1 for r in signers if r["months"][SEPTEMBER_INDEX][0] == STATE_PAID
    )
    sep_debt = sum(
        1
        for r in signers
        if r["months"][SEPTEMBER_INDEX][0] == STATE_DEBT and r["price"] > 0
    )
    hist_completed = 0
    hist_debt = 0
    nothing_owed = 0
    for row in roster:
        for index, (state, _amount) in enumerate(row["months"]):
            if index == SEPTEMBER_INDEX:
                if state == STATE_DEBT and row["price"] == 0:
                    nothing_owed += 1
                continue
            if state == STATE_PAID:
                hist_completed += 1
            elif state == STATE_DEBT:
                if row["price"] > 0:
                    hist_debt += 1
                else:
                    nothing_owed += 1

    return {
        "users": len(roster),
        "contracts": len(signers),
        "installments": len(signers),
        "enrolments": len(roster) + len(signers),
        "sep_completed": sep_completed,
        "sep_debt": sep_debt,
        "hist_completed": hist_completed,
        "hist_debt": hist_debt,
        "nothing_owed": nothing_owed,
        "payments": sep_completed + sep_debt + hist_completed + hist_debt,
    }


# ============================================
# GET-OR-CREATE HELPERS
# ============================================
# Every helper resolves against a pre-loaded index and mutates it. SessionLocal
# runs with autoflush=False, so a query issued mid-run would NOT see rows added
# earlier in the same run — index lookups are what keep this idempotent within
# a single pass, not just across runs. None of them commit.


def get_or_create_selection(db, index, company_id, name, counters):
    selection = index.get(name)
    if selection:
        counters.note("SELECTION", False, name, loud=True)
        return selection
    selection = Selection(company_id=company_id, name=name, is_active=True)
    db.add(selection)
    index[name] = selection
    counters.note("SELECTION", True, name, loud=True)
    return selection


def get_or_create_membership(db, index, company_id, selection, price, counters):
    key = (selection.id, PROGRAM, BILLING)
    membership = index.get(key)
    label = f"{selection.name} @ {price}"
    if membership:
        if Decimal(membership.price) != Decimal(price):
            counters.warn(
                f"membership {selection.name}: price is {membership.price} in the "
                f"database, {price} in the roster — left as is."
            )
        counters.note("MEMBERSHIP", False, label, loud=True)
        return membership
    membership = Membership(
        company_id=company_id,
        selection_id=selection.id,
        program=PROGRAM,
        billing_type=BILLING,
        price=Decimal(price),
        term_months=None,
        name=selection.name,
        is_active=True,
    )
    db.add(membership)
    index[key] = membership
    counters.note("MEMBERSHIP", True, label, loud=True)
    return membership


def get_or_create_season(db, index, company_id, spec, counters):
    name, start, end, is_current = spec
    season = index.get(name)
    if season:
        if bool(season.is_current) != is_current:
            counters.warn(
                f"season {name}: is_current is {season.is_current} in the database, "
                f"{is_current} in this script — left as is."
            )
        counters.note("SEASON", False, name, loud=True)
        return season
    season = Season(
        company_id=company_id,
        name=name,
        start_date=start,
        end_date=end,
        is_current=is_current,
    )
    db.add(season)
    index[name] = season
    counters.note("SEASON", True, name, loud=True)
    return season


def get_or_create_group(db, index, season, selection, counters):
    key = (season.id, selection.id)
    group = index.get(key)
    label = f"{season.name} / {selection.name}"
    if group:
        counters.note("GROUP", False, label, loud=True)
        return group
    group = Group(season_id=season.id, selection_id=selection.id)
    db.add(group)
    index[key] = group
    counters.note("GROUP", True, label, loud=True)
    return group


def get_or_create_user(db, index, emails, row, company_id, password_hash, counters):
    key = name_key(row["first_name"], row["last_name"], row["dob"])
    label = f"{row['first_name']} {row['last_name']} ({row['dob']})"
    user = index.get(key)
    if user:
        counters.note("USER", False, label)
        return user
    user = User(
        email=email_for(row["first_name"], row["last_name"], row["row_no"], emails),
        username=None,
        hashed_password=password_hash,
        first_name=row["first_name"],
        last_name=row["last_name"],
        phone_number=DEFAULT_PHONE,
        date_of_birth=row["dob"],
        company_id=company_id,
        roles=list(ROSTER_ROLES),
        is_active=True,
    )
    db.add(user)
    index[key] = user
    counters.note("USER", True, f"{label} -> {user.email}")
    return user


def get_or_create_wallet(db, index, owner_id, owner_type, name, counters):
    wallet = index.get(owner_id)
    if wallet:
        counters.note("WALLET", False, name)
        return wallet
    wallet = Wallet(owner_id=owner_id, owner_type=owner_type, name=name)
    db.add(wallet)
    index[owner_id] = wallet
    counters.note("WALLET", True, name)
    return wallet


def get_or_create_group_user(db, index, group, user, counters):
    key = (group.id, user.id)
    if key in index:
        counters.note("ENROLMENT", False)
        return None
    row = GroupUser(group_id=group.id, user_id=user.id)
    db.add(row)
    index.add(key)
    counters.note("ENROLMENT", True)
    return row


def get_or_create_contract(db, index, company_id, user, membership, amount, status,
                           counters):
    contract = index.get(user.id)
    label = f"{user.first_name} {user.last_name} @ {amount}"
    if contract:
        if Decimal(contract.amount) != Decimal(amount):
            counters.warn(
                f"contract for user {user.id}: amount is {contract.amount} in the "
                f"database, {amount} in the roster — left as is."
            )
        counters.note("CONTRACT", False, label)
        return contract
    contract = Contract(
        company_id=company_id,
        user_id=user.id,
        contract_type=ContractType.MEMBERSHIP,
        membership_id=membership.id,
        billing_type=BILLING,
        term_months=None,
        amount=Decimal(amount),
        start_date=CONTRACT_START,
        end_date=None,
        status=status,
        signed_at=datetime(
            CONTRACT_START.year, CONTRACT_START.month, CONTRACT_START.day,
            tzinfo=CLUB_TZ,
        ),
    )
    db.add(contract)
    index[user.id] = contract
    counters.note("CONTRACT", True, label)
    return contract


def get_or_create_installment(db, index, contract, period_start, period_end, counters):
    key = (contract.id, period_start)
    installment = index.get(key)
    if installment:
        counters.note("INSTALLMENT", False)
        return installment
    installment = ContractInstallment(
        contract_id=contract.id,
        period_start=period_start,
        period_end=period_end,
        due_date=period_start,
        amount=Decimal(contract.amount),
        waived=False,
    )
    db.add(installment)
    index[key] = installment
    counters.note("INSTALLMENT", True)
    return installment


def get_or_create_payment(db, index, sender_wallet_id, receiver_wallet_id, amount,
                          status, description, created_at, counters, tag,
                          payable_type=None, payable_id=None):
    """`payment` has no unique constraint, so the index IS the guard.

    (sender_wallet_id, description) is unique by construction: the description
    carries the month and the sender carries the user. `payable_id` is left out
    of the key on purpose — it is a consequence of the row, not its identity.
    """
    key = (sender_wallet_id, description)
    existing = index.get(key)
    if existing:
        if Decimal(existing.amount) != Decimal(amount) or existing.status != status:
            counters.warn(
                f"payment {description} for wallet {sender_wallet_id}: database has "
                f"({existing.amount}, {existing.status}), roster says "
                f"({amount}, {status}) — left as is."
            )
        counters.note(tag, False)
        return existing
    payment = Payment(
        sender_wallet_id=sender_wallet_id,
        receiver_wallet_id=receiver_wallet_id,
        payment_type=PaymentTypeCode.USER_MEMBERSHIP_FEE,
        amount=Decimal(amount),
        status=status,
        description=description,
        payable_type=payable_type,
        payable_id=payable_id,
        created_at=created_at,
    )
    db.add(payment)
    index[key] = payment
    counters.note(tag, True)
    return payment


# ============================================
# INDEX PRE-LOADERS — one SELECT each
# ============================================


def index_selections(db, company_id):
    return {
        s.name: s
        for s in db.query(Selection).filter(Selection.company_id == company_id).all()
    }


def index_memberships(db, company_id):
    return {
        (m.selection_id, m.program, m.billing_type): m
        for m in db.query(Membership).filter(Membership.company_id == company_id).all()
    }


def index_seasons(db, company_id):
    return {
        s.name: s
        for s in db.query(Season).filter(Season.company_id == company_id).all()
    }


def index_groups(db, season_ids):
    if not season_ids:
        return {}
    rows = db.query(Group).filter(Group.season_id.in_(season_ids)).all()
    return {(g.season_id, g.selection_id): g for g in rows}


def index_users(db, company_id):
    index = {}
    for user in db.query(User).filter(User.company_id == company_id).all():
        index.setdefault(name_key(user.first_name, user.last_name, user.date_of_birth),
                         user)
    return index


def index_all_emails(db):
    return {row[0] for row in db.query(User.email).all() if row[0]}


def index_user_wallets(db, owner_ids):
    if not owner_ids:
        return {}
    rows = (
        db.query(Wallet)
        .filter(
            Wallet.owner_type == WalletOwnerType.USER,
            Wallet.owner_id.in_(owner_ids),
        )
        .all()
    )
    return {w.owner_id: w for w in rows}


def index_group_users(db, group_ids):
    if not group_ids:
        return set()
    rows = db.query(GroupUser).filter(GroupUser.group_id.in_(group_ids)).all()
    return {(r.group_id, r.user_id) for r in rows}


def index_contracts(db, company_id):
    rows = (
        db.query(Contract)
        .filter(
            Contract.company_id == company_id,
            Contract.contract_type == ContractType.MEMBERSHIP,
            Contract.start_date == CONTRACT_START,
        )
        .all()
    )
    return {c.user_id: c for c in rows}


def index_installments(db, contract_ids):
    if not contract_ids:
        return {}
    rows = (
        db.query(ContractInstallment)
        .filter(ContractInstallment.contract_id.in_(contract_ids))
        .all()
    )
    return {(r.contract_id, r.period_start): r for r in rows}


def index_payments(db, receiver_wallet_id):
    rows = (
        db.query(Payment)
        .filter(Payment.receiver_wallet_id == receiver_wallet_id)
        .all()
    )
    return {
        (p.sender_wallet_id, p.description): p for p in rows if p.description
    }


# ============================================
# VERIFICATION — runs inside the transaction
# ============================================


def verify(db, club, academy, roster, expected, counters):
    problems = []

    def check(label, actual, wanted):
        if actual != wanted:
            problems.append(f"{label}: expected {wanted}, found {actual}")

    check("selections (academy)",
          db.query(Selection).filter(Selection.company_id == academy.id).count(),
          len(SELECTION_ORDER))
    check("memberships (academy)",
          db.query(Membership).filter(Membership.company_id == academy.id).count(),
          len(SELECTION_ORDER))
    check("current seasons (academy)",
          db.query(Season)
            .filter(Season.company_id == academy.id, Season.is_current.is_(True))
            .count(),
          1)

    season_ids = [
        s.id for s in db.query(Season).filter(Season.company_id == academy.id).all()
    ]
    check("groups (academy seasons)",
          db.query(Group).filter(Group.season_id.in_(season_ids)).count(),
          len(SELECTION_ORDER) * 2)

    contracts = db.query(Contract).filter(
        Contract.company_id == club.id,
        Contract.start_date == CONTRACT_START,
    ).all()
    check("contracts", len(contracts), expected["contracts"])
    check("non-ACTIVE contracts",
          sum(1 for c in contracts if ContractStatus(c.status) != ContractStatus.ACTIVE),
          0)
    check("open-ended contracts",
          sum(1 for c in contracts if c.end_date is None),
          expected["contracts"])

    contract_ids = [c.id for c in contracts]
    installments = (
        db.query(ContractInstallment)
        .filter(ContractInstallment.contract_id.in_(contract_ids))
        .all()
        if contract_ids else []
    )
    check("September installments", len(installments), expected["installments"])
    # As of this run: October and everything after it is the recurring job's.
    check("installments outside September 2026",
          sum(1 for i in installments if i.period_start != CONTRACT_START),
          0)

    # Every September installment has exactly one payment, except the zero-amount
    # ones, which must have none.
    by_id = {i.id: i for i in installments}
    linked = {}
    for payment in db.query(Payment).filter(
        Payment.payable_type == PayableType.CONTRACT_INSTALLMENT,
        Payment.payable_id.in_(list(by_id)) if by_id else False,
    ).all():
        linked[payment.payable_id] = linked.get(payment.payable_id, 0) + 1
    mismatched = [
        i.id for i in installments
        if (Decimal(i.amount) == 0) != (linked.get(i.id, 0) == 0)
        or linked.get(i.id, 0) > 1
    ]
    check("installments with the wrong payment count", len(mismatched), 0)

    payments = (
        db.query(Payment).filter(Payment.receiver_wallet_id == club.w_id).all()
    )
    check("payments", len(payments), expected["payments"])
    check("September COMPLETED",
          sum(1 for p in payments
              if p.payable_id is not None and p.status == PaymentStatus.COMPLETED),
          expected["sep_completed"])
    check("September DEBT",
          sum(1 for p in payments
              if p.payable_id is not None and p.status == PaymentStatus.DEBT),
          expected["sep_debt"])
    check("historical COMPLETED",
          sum(1 for p in payments
              if p.payable_id is None and p.status == PaymentStatus.COMPLETED),
          expected["hist_completed"])
    check("historical DEBT",
          sum(1 for p in payments
              if p.payable_id is None and p.status == PaymentStatus.DEBT),
          expected["hist_debt"])

    # The check that stands in for the missing unique constraint.
    keys = [(p.sender_wallet_id, p.description) for p in payments]
    check("duplicate payments", len(keys) - len(set(keys)), 0)

    check("club users without a wallet",
          db.query(User)
            .filter(User.company_id == club.id, User.w_id.is_(None))
            .count(),
          0)

    if problems:
        for problem in problems:
            print(f"  [FAIL       ] {problem}")
        raise AssertionError(f"{len(problems)} verification check(s) failed.")

    print("  [VERIFY     ] all checks passed.")


# ============================================
# THE SEED
# ============================================


def resolve_companies(db, club_id):
    club = db.get(Company, club_id)
    if club is None:
        raise SystemExit(f"No company with id {club_id}.")
    if club.academy_id is None:
        raise SystemExit(
            f"Company {club_id} ({club.name.strip()}) has no academy_id. Nothing "
            "academy-level (selection, season, group, membership) can be created "
            "until it belongs to an ACADEMY."
        )
    academy = db.get(Company, club.academy_id)
    if academy is None:
        raise SystemExit(f"Company {club_id} points at academy {club.academy_id}, "
                         "which does not exist.")
    if club.w_id is None:
        raise SystemExit(f"Company {club_id} has no wallet (w_id). Run seed_wallets.py.")
    return club, academy


def seed(path, club_id, dry_run, verbose):
    counters = Counters(verbose=verbose)
    current_row = None

    print(f"Importing the PK Taš011 roster from {os.path.basename(path)}")
    if dry_run:
        print("  *** DRY RUN — the transaction is rolled back at the end ***")

    roster = parse_roster(path)
    validate_roster(roster)
    expected = expected_counts(roster)
    print(f"  parsed {len(roster)} roster rows; "
          f"{expected['contracts']} have a September 2026 cell.")

    status = compute_contract_status(CONTRACT_START, None, club_today())
    if status != ContractStatus.ACTIVE:
        raise SystemExit(
            f"A contract starting {CONTRACT_START} computes as {status.value} today "
            f"({club_today()}), not ACTIVE. Seeding it would generate nothing. "
            "Run this on or after the start date."
        )

    db = SessionLocal()
    try:
        club, academy = resolve_companies(db, club_id)
        print(f"  club    id={club.id} {club.name.strip()!r} w_id={club.w_id}")
        print(f"  academy id={academy.id} {academy.name.strip()!r}")

        # --- academy: selections, memberships, seasons, groups ---
        selection_index = index_selections(db, academy.id)
        selections = {
            name: get_or_create_selection(db, selection_index, academy.id, name, counters)
            for name in SELECTION_ORDER
        }
        db.flush()

        membership_index = index_memberships(db, academy.id)
        memberships = {
            name: get_or_create_membership(
                db, membership_index, academy.id, selections[name],
                SELECTION_PRICES[name], counters,
            )
            for name in SELECTION_ORDER
        }
        db.flush()

        season_index = index_seasons(db, academy.id)
        season_prev = get_or_create_season(db, season_index, academy.id, SEASON_PREV,
                                           counters)
        season_curr = get_or_create_season(db, season_index, academy.id, SEASON_CURR,
                                           counters)
        db.flush()

        group_index = index_groups(db, [season_prev.id, season_curr.id])
        groups = {}
        for season in (season_prev, season_curr):
            for name in SELECTION_ORDER:
                groups[(season.id, name)] = get_or_create_group(
                    db, group_index, season, selections[name], counters
                )
        db.flush()

        # --- club: users ---
        password_hash = hash_password(DEFAULT_PASSWORD)
        user_index = index_users(db, club.id)
        emails = index_all_emails(db)
        users = {}
        for row in roster:
            current_row = row
            users[row["row_no"]] = get_or_create_user(
                db, user_index, emails, row, club.id, password_hash, counters
            )
        current_row = None
        db.flush()

        # --- club: wallets, then link them back onto the users ---
        wallet_index = index_user_wallets(db, [u.id for u in users.values()])
        wallets = {}
        for row in roster:
            user = users[row["row_no"]]
            wallets[row["row_no"]] = get_or_create_wallet(
                db, wallet_index, user.id, WalletOwnerType.USER,
                f"{user.first_name} {user.last_name}", counters,
            )
        db.flush()
        for row in roster:
            user = users[row["row_no"]]
            if user.w_id is None:
                user.w_id = wallets[row["row_no"]].id
        db.flush()

        # --- enrolment: everyone in 2025/2026, the September signers in 2026/2027 ---
        enrolment_index = index_group_users(db, [g.id for g in groups.values()])
        for row in roster:
            user = users[row["row_no"]]
            get_or_create_group_user(
                db, enrolment_index,
                groups[(season_prev.id, row["selection_name"])], user, counters,
            )
            if row["months"][SEPTEMBER_INDEX][0] != STATE_BLANK:
                get_or_create_group_user(
                    db, enrolment_index,
                    groups[(season_curr.id, row["selection_name"])], user, counters,
                )
        db.flush()

        # --- club: contracts for the September signers ---
        contract_index = index_contracts(db, club.id)
        contracts = {}
        for row in roster:
            if row["months"][SEPTEMBER_INDEX][0] == STATE_BLANK:
                continue
            current_row = row
            contracts[row["row_no"]] = get_or_create_contract(
                db, contract_index, club.id, users[row["row_no"]],
                memberships[row["selection_name"]], row["price"], status, counters,
            )
        current_row = None
        db.flush()

        # --- club: the September 2026 installment, and nothing later ---
        period_start, period_end = month_bounds(CONTRACT_START)
        installment_index = index_installments(db, [c.id for c in contracts.values()])
        installments = {}
        for row_no, contract in contracts.items():
            installments[row_no] = get_or_create_installment(
                db, installment_index, contract, period_start, period_end, counters
            )
        db.flush()

        # --- club: payments ---
        payment_index = index_payments(db, club.w_id)
        for row in roster:
            current_row = row
            sender = wallets[row["row_no"]].id
            for index, (state, amount) in enumerate(row["months"]):
                if state == STATE_BLANK:
                    continue
                year, month = MONTH_COLUMNS[index][1]
                description = DESCRIPTION_FMT.format(month, year)
                # Backdate to the month the money belongs to: created_at is the
                # payment list's sort key, and Postgres' now() would stamp all
                # 1282 rows with this transaction's start time.
                created_at = datetime(year, month, 5, 12, 0, tzinfo=CLUB_TZ)
                september = index == SEPTEMBER_INDEX

                if state == STATE_PAID:
                    # A green cell is money that moved, whatever the plan says —
                    # scholarship users included.
                    pay_amount = amount
                    pay_status = PaymentStatus.COMPLETED
                else:
                    # A red 0 means "owed but unpaid", and what is owed is the
                    # roster price. Nothing owed, nothing to record.
                    pay_amount = row["price"]
                    pay_status = PaymentStatus.DEBT
                    if pay_amount == 0:
                        counters.note("SKIPPED-ZERO", False)
                        continue

                installment = installments.get(row["row_no"]) if september else None
                get_or_create_payment(
                    db, payment_index, sender, club.w_id, pay_amount, pay_status,
                    description, created_at, counters,
                    tag="PAYMENT-SEP" if september else "PAYMENT-HIST",
                    payable_type=PayableType.CONTRACT_INSTALLMENT if installment else None,
                    payable_id=installment.id if installment else None,
                )
        current_row = None
        db.flush()

        verify(db, club, academy, roster, expected, counters)

        counters.print_summary()

        if dry_run:
            db.rollback()
            print("\n  DRY RUN — the transaction was rolled back, nothing was written.")
        else:
            db.commit()
            print("\n  Committed.")
    except Exception as exc:
        # The Postgres transaction may already be poisoned, so report from
        # in-memory state only — never query the session here.
        if current_row is not None:
            print(
                f"\n  Failed on roster row {current_row['row_no']} "
                f"({current_row['first_name']} {current_row['last_name']}): {exc}",
                file=sys.stderr,
            )
        db.rollback()
        raise
    finally:
        db.close()


def main() -> int:
    parser = argparse.ArgumentParser(
        description="One-off roster import for PK Taš011 from the season 25/26 xlsx."
    )
    parser.add_argument("--xlsx", default=XLSX_PATH, help="Path to the roster workbook.")
    parser.add_argument("--club-id", type=int, default=CLUB_ID,
                        help="The CLUB company id that owns the billing rows.")
    parser.add_argument("--dry-run", action="store_true",
                        help="Run everything, assert, then roll back.")
    parser.add_argument("--verbose", action="store_true",
                        help="Print a line per user, contract, installment and payment.")
    args = parser.parse_args()

    seed(args.xlsx, args.club_id, args.dry_run, args.verbose)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
